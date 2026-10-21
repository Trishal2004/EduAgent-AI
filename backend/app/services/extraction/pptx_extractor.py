from pptx import Presentation
from typing import List
from uuid import UUID
from app.services.extraction.base import BaseExtractor
from app.DTOs.document_chunk import DocumentChunkCreate

class PPTXExtractor(BaseExtractor):
    def extract(self, file_path: str, document_id: UUID) -> List[DocumentChunkCreate]:
        chunks = []
        prs = Presentation(file_path)
        
        for i, slide in enumerate(prs.slides):
            slide_text = []
            slide_title = f"Slide {i + 1}"
            
            # Try to find a title shape
            if slide.shapes.title and slide.shapes.title.has_text_frame:
                slide_title = slide.shapes.title.text.strip()
                
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text:
                    # Don't duplicate the title in the body text if possible, 
                    # but for simplicity we'll just grab all text
                    slide_text.append(shape.text.strip())
            
            text_content = "\n".join(slide_text).strip()
            if text_content:
                chunks.append(
                    DocumentChunkCreate(
                        document_id=document_id,
                        page_number=i + 1,
                        section_title=slide_title,
                        text_content=text_content,
                        extracted_metadata={"source_type": "pptx_slide"}
                    )
                )
                
        return chunks
